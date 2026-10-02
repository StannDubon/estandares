import os, io, re, secrets, datetime as dt
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, abort, session
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
import openpyxl

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'cambiar-en-produccion')
app.config['SQLALCHEMY_DATABASE_URI'] = re.sub(r'^postgres(ql)?://', 'postgresql+psycopg2://', os.environ.get('DATABASE_URL', 'sqlite:///local.db'))
db = SQLAlchemy(app)
lm = LoginManager(app); lm.login_view = 'login'

# (campo BD, columna en el Excel, etiqueta)
FIELDS = [('pae_pri','B','PAE priorizados'),('pae_real','C','PAE realizados'),
 ('pap_total','E','Muestras PAP/VPH procesadas (solo enfermería)'),('pap_sat','F','PAP/VPH satisfactorio'),('pap_insat','G','PAP/VPH insatisfactorio'),
 ('enc_total','I','Usuarios encuestados'),('enc_sat','J','Usuarios satisfechos'),('enc_insat','K','Usuarios insatisfechos'),
 ('tam_total','M','Tamizajes tomados'),('tam_ok','N','Tamizajes satisfactorios'),('tam_rech','O','Tamizajes rechazados'),
 ('ven_total','Q','Usuarios con vena canalizada'),('ven_fleb','R','Usuarios con flebitis'),('caidas','S','Caídas de usuarios')]
KEYS = [f[0] for f in FIELDS]
# después de qué campo va un %: (etiqueta, numerador, denominador, col Excel, col num, col den)
PCT = {'pae_real':('% PAE realizados','pae_real','pae_pri','D','C','B'),
       'pap_insat':('% lectura satisfactoria','pap_sat','pap_total','H','F','E'),
       'enc_insat':('% usuarios satisfechos','enc_sat','enc_total','L','J','I'),
       'tam_rech':('% muestras aceptadas','tam_ok','tam_total','P','N','M')}
MESES = ['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE']

class Region(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(80), unique=True, nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)

class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    full_name = db.Column(db.String(120), nullable=False)
    username = db.Column(db.String(50), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    is_admin = db.Column(db.Boolean, default=False, nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)
    region_id = db.Column(db.Integer, db.ForeignKey('region.id'))
    region = db.relationship('Region')
    @property
    def is_active(self): return self.active

class Entry(db.Model):
    __tablename__ = 'entries'
    id = db.Column(db.Integer, primary_key=True)
    region_id = db.Column(db.Integer, db.ForeignKey('region.id'), nullable=False)
    year = db.Column(db.Integer, nullable=False)
    month = db.Column(db.Integer, nullable=False)
    updated_by = db.Column(db.Integer, db.ForeignKey('user.id'))
    updated_at = db.Column(db.DateTime)
    __table_args__ = (db.UniqueConstraint('region_id', 'year', 'month'),)
for _k in KEYS: setattr(Entry, _k, db.Column(db.Integer))

class Audit(db.Model):  # historial: nunca se borra
    id = db.Column(db.Integer, primary_key=True)
    entry_id = db.Column(db.Integer, db.ForeignKey('entries.id'), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    field = db.Column(db.String(30)); old = db.Column(db.Integer); new = db.Column(db.Integer)
    at = db.Column(db.DateTime, default=dt.datetime.utcnow)
    entry = db.relationship('Entry'); user = db.relationship('User')

class Setting(db.Model):  # periodo habilitado para los usuarios
    key = db.Column(db.String(30), primary_key=True); value = db.Column(db.String(30))

def period():
    y, m, t = db.session.get(Setting, 'year'), db.session.get(Setting, 'month'), dt.date.today()
    return (int(y.value) if y else t.year, int(m.value) if m else t.month)

def pct(a, b): return round(a * 100 / b, 1) if a is not None and b else ''
app.jinja_env.globals.update(pct=pct, F=FIELDS, P=PCT, MESES=MESES)

@lm.user_loader
def load_user(i): return db.session.get(User, int(i))

@app.before_request
def csrf():
    if request.method == 'POST' and (not session.get('_t') or request.form.get('_t') != session['_t']): abort(400)

@app.context_processor
def tok():
    session.setdefault('_t', secrets.token_hex(16)); return {'T': session['_t']}

def admin_only():
    if not current_user.is_admin: abort(403)

def setup():
    db.create_all()
    pg = db.engine.dialect.name == 'postgresql'
    with db.engine.begin() as c:
        for t in ('entries', 'audit'):  # la BD misma impide borrar
            if pg:
                c.execute(text("CREATE OR REPLACE FUNCTION no_delete() RETURNS trigger AS $$ BEGIN RAISE EXCEPTION 'No se permite borrar datos'; END; $$ LANGUAGE plpgsql"))
                c.execute(text(f"DROP TRIGGER IF EXISTS no_del ON {t}"))
                c.execute(text(f"CREATE TRIGGER no_del BEFORE DELETE ON {t} FOR EACH ROW EXECUTE FUNCTION no_delete()"))
            else:
                c.execute(text(f"CREATE TRIGGER IF NOT EXISTS no_del_{t} BEFORE DELETE ON {t} BEGIN SELECT RAISE(ABORT,'No se permite borrar datos'); END"))
    if not Region.query.count():
        for n in ['San Ramon','Citala','Las Pilas','San Ignacio','Granadillas','San Jose','Horcones','La Palma']: db.session.add(Region(name=n))
    if not User.query.filter_by(is_admin=True).count():
        db.session.add(User(full_name='Administrador', username=os.environ.get('ADMIN_USER', 'admin'), is_admin=True,
                            password_hash=generate_password_hash(os.environ.get('ADMIN_PASSWORD', 'cambiar123'))))
    db.session.commit()
with app.app_context(): setup()

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        u = User.query.filter_by(username=request.form['username'].strip()).first()
        if u and u.active and check_password_hash(u.password_hash, request.form['password']):
            login_user(u); return redirect(url_for('home'))
        flash('Usuario o contraseña incorrectos')
    return render_template('login.html')

@app.route('/logout', methods=['POST'])
def logout(): logout_user(); return redirect(url_for('login'))

@app.route('/')
@login_required
def home():
    if current_user.is_admin: return redirect(url_for('admin'))
    if not current_user.region_id: return 'Sin región asignada. Contacte al administrador.'
    return redirect(url_for('region', rid=current_user.region_id))

@app.route('/region/<int:rid>', methods=['GET', 'POST'])
@login_required
def region(rid):
    r = db.get_or_404(Region, rid)
    if not current_user.is_admin and (current_user.region_id != rid or not r.active): abort(403)
    py, pm = period()
    if current_user.is_admin: year = request.values.get('year', type=int) or py; editable = set(range(1, 13))
    else: year = py; editable = {pm}  # el usuario solo edita el mes que habilita el admin
    ents = {e.month: e for e in Entry.query.filter_by(region_id=rid, year=year)}
    if request.method == 'POST':
        for m in editable:
            for k in KEYS:
                raw = request.form.get(f'{m}_{k}', '').strip()
                try: v = int(raw) if raw else None
                except ValueError: flash(f'Valor inválido en {MESES[m-1]}'); continue
                if v is not None and v < 0: continue
                e = ents.get(m); old = getattr(e, k) if e else None
                if v == old: continue
                if not e:
                    e = ents[m] = Entry(region_id=rid, year=year, month=m); db.session.add(e); db.session.flush()
                setattr(e, k, v); e.updated_by = current_user.id; e.updated_at = dt.datetime.utcnow()
                db.session.add(Audit(entry_id=e.id, user_id=current_user.id, field=k, old=old, new=v))
        db.session.commit(); flash('Guardado'); return redirect(url_for('region', rid=rid, year=year))
    vals = {m: {k: getattr(ents[m], k) if m in ents else None for k in KEYS} for m in range(1, 13)}
    return render_template('region.html', r=r, year=year, vals=vals, editable=editable, pm=pm, py=py)

@app.route('/admin')
@login_required
def admin():
    admin_only()
    return render_template('admin.html', period=period(), users=User.query.order_by(User.id).all(), regions=Region.query.order_by(Region.id).all())

@app.post('/admin/region')
@login_required
def admin_region():
    admin_only(); i = request.form.get('id')
    r = db.session.get(Region, int(i)) if i else Region()
    r.name = request.form['name'].strip(); r.active = bool(request.form.get('active')) if i else True
    db.session.add(r)
    try: db.session.commit(); flash('Región guardada')
    except IntegrityError: db.session.rollback(); flash('Ese nombre ya existe')
    return redirect(url_for('admin'))

@app.post('/admin/user')
@login_required
def admin_user():
    admin_only(); i = request.form.get('id')
    u = db.session.get(User, int(i)) if i else User()
    u.full_name = request.form['full_name'].strip(); u.username = request.form['username'].strip()
    u.region_id = int(request.form['region_id']) if request.form.get('region_id') else None
    u.active = bool(request.form.get('active')) if i else True
    u.is_admin = bool(request.form.get('is_admin'))
    if request.form.get('password'): u.password_hash = generate_password_hash(request.form['password'])
    elif not i: flash('Falta contraseña'); return redirect(url_for('admin'))
    db.session.add(u)
    try: db.session.commit(); flash('Usuario guardado')
    except IntegrityError: db.session.rollback(); flash('Ese usuario ya existe')
    return redirect(url_for('admin'))

@app.post('/admin/period')
@login_required
def admin_period():
    admin_only()
    for k in ('year', 'month'):
        v = db.session.get(Setting, k) or Setting(key=k); v.value = str(int(request.form[k])); db.session.add(v)
    db.session.commit(); flash('Mes habilitado actualizado'); return redirect(url_for('admin'))

@app.route('/audit')
@login_required
def audit():
    admin_only()
    return render_template('audit.html', rows=Audit.query.order_by(Audit.id.desc()).limit(300).all(), MESES=MESES)

def fill(ws, d, year, title):
    ws['A5'] = str(year); ws['A6'] = 'NOMBRE DE LA UNIDAD DE SALUD: ' + title
    for r in range(9, 21):
        for c in range(2, 20): ws.cell(r, c).value = None
        for k, col, _ in FIELDS:
            if k in d.get(r - 8, {}): ws[f'{col}{r}'] = d[r - 8][k]
    for _, col, _ in FIELDS: ws[f'{col}21'] = f'=SUM({col}9:{col}20)'
    for r in range(9, 22):
        for _, _, _, pc, n, dn in PCT.values(): ws[f'{pc}{r}'] = f'=IF({dn}{r}>0,ROUND({n}{r}/{dn}{r}*100,1),"")'

def agg(entries):
    d = {}
    for e in entries:
        for k in KEYS:
            v = getattr(e, k)
            if v is not None: x = d.setdefault(e.month, {}); x[k] = x.get(k, 0) + v
    return d

@app.route('/export')
@login_required
def export():
    admin_only(); year = request.args.get('year', type=int) or dt.date.today().year
    wb = openpyxl.load_workbook(os.path.join(os.path.dirname(__file__), 'plantilla.xlsx')); base = wb.worksheets[0]
    regs = Region.query.order_by(Region.id).all()
    sheets = []
    for r in regs:
        ws = wb.copy_worksheet(base); ws.title = re.sub(r'[\[\]:*?/\\]', '', r.name)[:31]; sheets.append((ws, r))
    for ws, r in sheets: fill(ws, agg(Entry.query.filter_by(region_id=r.id, year=year)), year, r.name.upper())
    fill(base, agg(Entry.query.filter_by(year=year)), year, 'CONSOLIDADO (TODAS LAS REGIONES)'); base.title = 'CONSOLIDADO'
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'ESTANDARES_ENFERMERIA_{year}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
